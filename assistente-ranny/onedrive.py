"""
🔐 OneDrive Integration - OAuth2 com Microsoft Graph API
Assistente Ranny V3

Implementa autenticação OAuth2 e acesso ao OneDrive via Microsoft Graph API.
Requirements: 7.1, 7.2
"""

import os
import logging
from datetime import datetime, timedelta
from typing import Optional, List
from urllib.parse import urlencode
import httpx

from config import (
    MICROSOFT_CLIENT_ID,
    MICROSOFT_CLIENT_SECRET,
    MICROSOFT_REDIRECT_URI,
    MICROSOFT_SCOPES
)
from database import (
    get_oauth_token,
    save_oauth_token,
    update_oauth_token,
    delete_oauth_token
)

# Logging
logger = logging.getLogger(__name__)

# Microsoft OAuth2 endpoints
MICROSOFT_AUTH_URL = "https://login.microsoftonline.com/common/oauth2/v2.0/authorize"
MICROSOFT_TOKEN_URL = "https://login.microsoftonline.com/common/oauth2/v2.0/token"
GRAPH_API_URL = "https://graph.microsoft.com/v1.0"


class OneDriveAuth:
    """
    Gerencia autenticação OAuth2 com Microsoft Graph API.
    
    Fluxo de autenticação:
    1. Usuário solicita conexão com OneDrive
    2. Bot gera URL de autorização e envia para usuário
    3. Usuário autoriza e é redirecionado com código
    4. Bot troca código por tokens
    5. Tokens são salvos no banco de dados local
    6. Access token é renovado automaticamente quando expira
    """
    
    def __init__(self):
        self.client_id = MICROSOFT_CLIENT_ID
        self.client_secret = MICROSOFT_CLIENT_SECRET
        self.redirect_uri = MICROSOFT_REDIRECT_URI
        self.scopes = MICROSOFT_SCOPES
        self.provider = 'microsoft'
    
    def is_configured(self) -> bool:
        """Verifica se as credenciais Microsoft estão configuradas"""
        return bool(self.client_id and self.client_secret)
    
    def get_auth_url(self, state: str = None) -> str:
        """
        Gera URL de autorização OAuth2.
        
        Args:
            state: Valor opcional para prevenir CSRF
        
        Returns:
            URL para redirecionar o usuário para autorização
        """
        params = {
            'client_id': self.client_id,
            'response_type': 'code',
            'redirect_uri': self.redirect_uri,
            'scope': ' '.join(self.scopes),
            'response_mode': 'query',
        }
        
        if state:
            params['state'] = state
        
        return f"{MICROSOFT_AUTH_URL}?{urlencode(params)}"
    
    async def exchange_code(self, code: str) -> bool:
        """
        Troca código de autorização por tokens.
        
        Args:
            code: Código recebido após autorização do usuário
        
        Returns:
            True se tokens foram obtidos e salvos com sucesso
        """
        if not self.is_configured():
            logger.error("Microsoft OAuth não configurado")
            return False
        
        data = {
            'client_id': self.client_id,
            'client_secret': self.client_secret,
            'code': code,
            'redirect_uri': self.redirect_uri,
            'grant_type': 'authorization_code',
        }
        
        try:
            async with httpx.AsyncClient() as client:
                response = await client.post(
                    MICROSOFT_TOKEN_URL,
                    data=data,
                    headers={'Content-Type': 'application/x-www-form-urlencoded'}
                )
                
                if response.status_code != 200:
                    logger.error(f"Erro ao trocar código: {response.status_code} - {response.text}")
                    return False
                
                tokens = response.json()
                
                # Calcula data de expiração
                expires_in = tokens.get('expires_in', 3600)  # Default 1 hora
                expires_at = datetime.now() + timedelta(seconds=expires_in)
                
                # Salva tokens no banco
                save_oauth_token(
                    provider=self.provider,
                    access_token=tokens.get('access_token'),
                    refresh_token=tokens.get('refresh_token'),
                    expires_at=expires_at,
                    scope=tokens.get('scope')
                )
                
                logger.info("Tokens Microsoft salvos com sucesso")
                return True
                
        except Exception as e:
            logger.error(f"Erro ao trocar código por tokens: {e}")
            return False
    
    async def refresh_access_token(self) -> bool:
        """
        Renova o access token usando o refresh token.
        
        Returns:
            True se token foi renovado com sucesso
        """
        token_data = get_oauth_token(self.provider)
        
        if not token_data or not token_data.get('refresh_token'):
            logger.warning("Sem refresh token disponível")
            return False
        
        data = {
            'client_id': self.client_id,
            'client_secret': self.client_secret,
            'refresh_token': token_data['refresh_token'],
            'grant_type': 'refresh_token',
        }
        
        try:
            async with httpx.AsyncClient() as client:
                response = await client.post(
                    MICROSOFT_TOKEN_URL,
                    data=data,
                    headers={'Content-Type': 'application/x-www-form-urlencoded'}
                )
                
                if response.status_code != 200:
                    logger.error(f"Erro ao renovar token: {response.status_code} - {response.text}")
                    # Se refresh falhou, pode ser que o token foi revogado
                    if response.status_code == 400:
                        delete_oauth_token(self.provider)
                    return False
                
                tokens = response.json()
                
                # Calcula nova data de expiração
                expires_in = tokens.get('expires_in', 3600)
                expires_at = datetime.now() + timedelta(seconds=expires_in)
                
                # Se veio novo refresh_token, salva tudo
                if tokens.get('refresh_token'):
                    save_oauth_token(
                        provider=self.provider,
                        access_token=tokens.get('access_token'),
                        refresh_token=tokens.get('refresh_token'),
                        expires_at=expires_at,
                        scope=tokens.get('scope')
                    )
                else:
                    # Apenas atualiza access_token
                    update_oauth_token(
                        provider=self.provider,
                        access_token=tokens.get('access_token'),
                        expires_at=expires_at
                    )
                
                logger.info("Token Microsoft renovado com sucesso")
                return True
                
        except Exception as e:
            logger.error(f"Erro ao renovar token: {e}")
            return False
    
    async def get_valid_token(self) -> Optional[str]:
        """
        Retorna um access token válido.
        
        Se o token atual expirou, tenta renovar automaticamente.
        
        Returns:
            Access token válido ou None se não disponível
        """
        token_data = get_oauth_token(self.provider)
        
        if not token_data:
            return None
        
        # Verifica se token expirou (com margem de 5 minutos)
        expires_at_str = token_data.get('expires_at')
        if expires_at_str:
            try:
                # Parse da data de expiração
                if isinstance(expires_at_str, str):
                    # Remove timezone info se presente
                    expires_at_str = expires_at_str.replace('Z', '').split('+')[0]
                    expires_at = datetime.fromisoformat(expires_at_str)
                else:
                    expires_at = expires_at_str
                
                # Se expira em menos de 5 minutos, renova
                if expires_at <= datetime.now() + timedelta(minutes=5):
                    logger.info("Token expirando, renovando...")
                    if not await self.refresh_access_token():
                        return None
                    # Busca token atualizado
                    token_data = get_oauth_token(self.provider)
                    
            except Exception as e:
                logger.error(f"Erro ao verificar expiração: {e}")
        
        return token_data.get('access_token') if token_data else None
    
    async def is_connected(self) -> bool:
        """
        Verifica se há conexão válida com OneDrive.
        
        Property 14: OneDrive desconectado retorna mensagem apropriada
        Requirements: 7.2
        
        Returns:
            True se conectado e token válido
        """
        token = await self.get_valid_token()
        
        if not token:
            return False
        
        # Testa conexão fazendo uma chamada simples à API
        try:
            async with httpx.AsyncClient() as client:
                response = await client.get(
                    f"{GRAPH_API_URL}/me/drive",
                    headers={'Authorization': f'Bearer {token}'},
                    timeout=10.0
                )
                return response.status_code == 200
                
        except Exception as e:
            logger.error(f"Erro ao verificar conexão OneDrive: {e}")
            return False
    
    async def get_connection_status(self) -> dict:
        """
        Retorna status de conexão com mensagem apropriada.
        
        Property 14: OneDrive desconectado retorna mensagem apropriada
        Requirements: 7.2
        
        Returns:
            dict com 'connected' (bool) e 'message' (str)
        """
        # Verifica se há token salvo
        token_data = get_oauth_token(self.provider)
        
        if not token_data:
            return {
                'connected': False,
                'message': "OneDrive não está conectado. Use o link de autorização para conectar."
            }
        
        # Verifica se token é válido fazendo chamada à API
        token = await self.get_valid_token()
        
        if not token:
            return {
                'connected': False,
                'message': "Não consegui validar a conexão. Tenta reconectar o OneDrive."
            }
        
        # Testa conexão real com a API
        try:
            async with httpx.AsyncClient() as client:
                response = await client.get(
                    f"{GRAPH_API_URL}/me/drive",
                    headers={'Authorization': f'Bearer {token}'},
                    timeout=10.0
                )
                
                if response.status_code == 200:
                    return {
                        'connected': True,
                        'message': "OneDrive conectado e funcionando! 🟢"
                    }
                elif response.status_code == 401:
                    return {
                        'connected': False,
                        'message': "Autorização expirou. Precisa reconectar o OneDrive."
                    }
                else:
                    return {
                        'connected': False,
                        'message': "Não consegui acessar o OneDrive. Tenta de novo! 😅"
                    }
                    
        except httpx.TimeoutException:
            return {
                'connected': False,
                'message': "OneDrive demorou pra responder. Tenta de novo! ⏱️"
            }
        except Exception as e:
            logger.error(f"Erro ao verificar conexão OneDrive: {e}")
            return {
                'connected': False,
                'message': "Tive um probleminha ao acessar o OneDrive. Tenta de novo! 😅"
            }
    
    def disconnect(self) -> bool:
        """
        Desconecta do OneDrive removendo tokens.
        
        Returns:
            True se tokens foram removidos
        """
        return delete_oauth_token(self.provider)


class OneDriveClient:
    """
    Cliente para operações no OneDrive via Microsoft Graph API.
    
    Requirements: 7.1, 7.3, 7.4
    """
    
    # Pasta padrão para sincronização (pode ser configurada via env)
    DEFAULT_SYNC_FOLDER = os.getenv('ONEDRIVE_SYNC_FOLDER', '')
    
    def __init__(self):
        self.auth = OneDriveAuth()
        self.sync_folder = self.DEFAULT_SYNC_FOLDER
    
    async def is_connected(self) -> bool:
        """Verifica se está conectado ao OneDrive"""
        return await self.auth.is_connected()
    
    async def get_connection_status(self) -> dict:
        """
        Retorna status de conexão com mensagem apropriada.
        
        Property 14: OneDrive desconectado retorna mensagem apropriada
        Requirements: 7.2
        
        Returns:
            dict com 'connected' (bool) e 'message' (str)
        """
        return await self.auth.get_connection_status()
    
    async def _get_headers(self) -> Optional[dict]:
        """Retorna headers com token de autorização"""
        token = await self.auth.get_valid_token()
        if not token:
            return None
        return {'Authorization': f'Bearer {token}'}
    
    async def search_files(self, query: str, limit: int = 10, folder_path: str = None, 
                          file_types: List[str] = None) -> List[dict]:
        """
        Busca arquivos no OneDrive.
        
        Busca na pasta sincronizada (se configurada) ou na raiz do OneDrive.
        Suporta filtro por tipo de arquivo.
        
        Args:
            query: Termo de busca (nome do arquivo, parte do nome, etc)
            limit: Número máximo de resultados (padrão: 10)
            folder_path: Caminho da pasta para buscar (opcional, usa sync_folder se não especificado)
            file_types: Lista de extensões para filtrar (ex: ['pdf', 'docx'])
        
        Returns:
            Lista de arquivos encontrados com metadados:
            - id: ID único do arquivo no OneDrive
            - name: Nome do arquivo
            - size: Tamanho em bytes
            - size_formatted: Tamanho formatado (KB, MB, etc)
            - web_url: URL para abrir no navegador
            - created: Data de criação
            - modified: Data de modificação
            - is_folder: Se é pasta
            - download_url: URL direta para download
            - path: Caminho completo do arquivo
            - extension: Extensão do arquivo
            - mime_type: Tipo MIME do arquivo
        
        Requirements: 7.1
        """
        headers = await self._get_headers()
        if not headers:
            logger.warning("OneDrive não conectado - não é possível buscar arquivos")
            return []
        
        try:
            async with httpx.AsyncClient() as client:
                # Determina o caminho de busca
                search_path = folder_path or self.sync_folder
                
                if search_path:
                    # Busca em pasta específica
                    # Escapa caracteres especiais no caminho
                    encoded_path = search_path.replace("'", "''")
                    search_url = f"{GRAPH_API_URL}/me/drive/root:/{encoded_path}:/search(q='{query}')"
                else:
                    # Busca na raiz do OneDrive
                    search_url = f"{GRAPH_API_URL}/me/drive/root/search(q='{query}')"
                
                # Parâmetros da busca
                params = {
                    '$top': limit,
                    '$select': 'id,name,size,webUrl,createdDateTime,lastModifiedDateTime,folder,file,parentReference,@microsoft.graph.downloadUrl'
                }
                
                response = await client.get(
                    search_url,
                    headers=headers,
                    params=params,
                    timeout=30.0
                )
                
                if response.status_code == 404:
                    # Pasta não encontrada - tenta buscar na raiz
                    logger.warning(f"Pasta '{search_path}' não encontrada, buscando na raiz")
                    search_url = f"{GRAPH_API_URL}/me/drive/root/search(q='{query}')"
                    response = await client.get(
                        search_url,
                        headers=headers,
                        params=params,
                        timeout=30.0
                    )
                
                if response.status_code != 200:
                    logger.error(f"Erro na busca OneDrive: {response.status_code} - {response.text}")
                    return []
                
                data = response.json()
                files = []
                
                for item in data.get('value', []):
                    # Extrai informações do arquivo
                    file_info = self._parse_file_item(item)
                    
                    # Filtra por tipo de arquivo se especificado
                    if file_types:
                        ext = file_info.get('extension', '').lower()
                        if ext not in [t.lower().lstrip('.') for t in file_types]:
                            continue
                    
                    files.append(file_info)
                
                logger.info(f"Busca OneDrive '{query}': {len(files)} arquivo(s) encontrado(s)")
                return files
                
        except httpx.TimeoutException:
            logger.error("Timeout ao buscar arquivos no OneDrive")
            return []
        except Exception as e:
            logger.error(f"Erro ao buscar arquivos no OneDrive: {e}")
            return []
    
    def _parse_file_item(self, item: dict) -> dict:
        """
        Converte item da API do Graph para formato padronizado.
        
        Args:
            item: Item retornado pela API do Microsoft Graph
        
        Returns:
            dict com informações formatadas do arquivo
        """
        name = item.get('name', '')
        size = item.get('size', 0)
        
        # Extrai extensão do nome
        extension = ''
        if '.' in name and not item.get('folder'):
            extension = name.rsplit('.', 1)[-1].lower()
        
        # Formata tamanho
        size_formatted = self._format_size(size)
        
        # Extrai caminho completo
        parent_ref = item.get('parentReference', {})
        parent_path = parent_ref.get('path', '')
        # Remove o prefixo /drive/root: se presente
        if parent_path.startswith('/drive/root:'):
            parent_path = parent_path[12:]
        full_path = f"{parent_path}/{name}" if parent_path else name
        
        # Extrai tipo MIME
        file_info = item.get('file', {})
        mime_type = file_info.get('mimeType', '')
        
        return {
            'id': item.get('id'),
            'name': name,
            'size': size,
            'size_formatted': size_formatted,
            'web_url': item.get('webUrl'),
            'created': item.get('createdDateTime'),
            'modified': item.get('lastModifiedDateTime'),
            'is_folder': 'folder' in item,
            'download_url': item.get('@microsoft.graph.downloadUrl'),
            'path': full_path,
            'extension': extension,
            'mime_type': mime_type,
            'parent_id': parent_ref.get('id'),
            'drive_id': parent_ref.get('driveId')
        }
    
    def _format_size(self, size_bytes: int) -> str:
        """
        Formata tamanho em bytes para formato legível.
        
        Args:
            size_bytes: Tamanho em bytes
        
        Returns:
            String formatada (ex: "1.5 MB", "256 KB")
        """
        if size_bytes == 0:
            return "0 B"
        
        units = ['B', 'KB', 'MB', 'GB', 'TB']
        unit_index = 0
        size = float(size_bytes)
        
        while size >= 1024 and unit_index < len(units) - 1:
            size /= 1024
            unit_index += 1
        
        if unit_index == 0:
            return f"{int(size)} {units[unit_index]}"
        return f"{size:.1f} {units[unit_index]}"
    
    async def search_in_sync_folder(self, query: str, limit: int = 10, 
                                    file_types: List[str] = None) -> List[dict]:
        """
        Busca arquivos especificamente na pasta sincronizada.
        
        Atalho para search_files usando a pasta de sincronização configurada.
        
        Args:
            query: Termo de busca
            limit: Número máximo de resultados
            file_types: Lista de extensões para filtrar
        
        Returns:
            Lista de arquivos encontrados
        
        Requirements: 7.1
        """
        return await self.search_files(
            query=query,
            limit=limit,
            folder_path=self.sync_folder,
            file_types=file_types
        )
    
    async def list_folder(self, folder_path: str = None, limit: int = 50) -> List[dict]:
        """
        Lista arquivos em uma pasta específica.
        
        Args:
            folder_path: Caminho da pasta (usa sync_folder se não especificado)
            limit: Número máximo de itens
        
        Returns:
            Lista de arquivos e pastas
        """
        headers = await self._get_headers()
        if not headers:
            return []
        
        try:
            async with httpx.AsyncClient() as client:
                path = folder_path or self.sync_folder
                
                if path:
                    encoded_path = path.replace("'", "''")
                    url = f"{GRAPH_API_URL}/me/drive/root:/{encoded_path}:/children"
                else:
                    url = f"{GRAPH_API_URL}/me/drive/root/children"
                
                response = await client.get(
                    url,
                    headers=headers,
                    params={'$top': limit},
                    timeout=30.0
                )
                
                if response.status_code != 200:
                    logger.error(f"Erro ao listar pasta: {response.status_code}")
                    return []
                
                data = response.json()
                return [self._parse_file_item(item) for item in data.get('value', [])]
                
        except Exception as e:
            logger.error(f"Erro ao listar pasta: {e}")
            return []
    
    async def download_file(self, file_id: str) -> Optional[bytes]:
        """
        Baixa um arquivo do OneDrive.
        
        Args:
            file_id: ID do arquivo no OneDrive
        
        Returns:
            Conteúdo do arquivo em bytes ou None
        
        Requirements: 7.4
        """
        headers = await self._get_headers()
        if not headers:
            logger.warning("OneDrive não conectado - não é possível baixar arquivo")
            return None
        
        try:
            async with httpx.AsyncClient() as client:
                # Primeiro, obtém a URL de download
                response = await client.get(
                    f"{GRAPH_API_URL}/me/drive/items/{file_id}",
                    headers=headers,
                    params={'$select': 'id,name,@microsoft.graph.downloadUrl'},
                    timeout=10.0
                )
                
                if response.status_code != 200:
                    logger.error(f"Erro ao obter arquivo: {response.status_code}")
                    return None
                
                item = response.json()
                download_url = item.get('@microsoft.graph.downloadUrl')
                
                if not download_url:
                    logger.error("URL de download não disponível")
                    return None
                
                # Baixa o arquivo
                download_response = await client.get(
                    download_url,
                    timeout=60.0  # Timeout maior para downloads
                )
                
                if download_response.status_code != 200:
                    logger.error(f"Erro no download: {download_response.status_code}")
                    return None
                
                logger.info(f"Arquivo '{item.get('name')}' baixado com sucesso")
                return download_response.content
                
        except httpx.TimeoutException:
            logger.error("Timeout ao baixar arquivo do OneDrive")
            return None
        except Exception as e:
            logger.error(f"Erro ao baixar arquivo: {e}")
            return None
    
    async def get_file_info(self, file_id: str) -> Optional[dict]:
        """
        Obtém informações detalhadas de um arquivo.
        
        Args:
            file_id: ID do arquivo no OneDrive
        
        Returns:
            dict com informações do arquivo ou None
        """
        headers = await self._get_headers()
        if not headers:
            return None
        
        try:
            async with httpx.AsyncClient() as client:
                response = await client.get(
                    f"{GRAPH_API_URL}/me/drive/items/{file_id}",
                    headers=headers,
                    timeout=10.0
                )
                
                if response.status_code != 200:
                    return None
                
                return self._parse_file_item(response.json())
                
        except Exception as e:
            logger.error(f"Erro ao obter info do arquivo: {e}")
            return None
    
    async def get_recent_files(self, limit: int = 10) -> List[dict]:
        """
        Lista arquivos recentes do OneDrive.
        
        Args:
            limit: Número máximo de arquivos
        
        Returns:
            Lista de arquivos recentes com metadados completos
        """
        headers = await self._get_headers()
        if not headers:
            return []
        
        try:
            async with httpx.AsyncClient() as client:
                response = await client.get(
                    f"{GRAPH_API_URL}/me/drive/recent",
                    headers=headers,
                    params={'$top': limit},
                    timeout=30.0
                )
                
                if response.status_code != 200:
                    logger.error(f"Erro ao listar recentes: {response.status_code}")
                    return []
                
                data = response.json()
                return [self._parse_file_item(item) for item in data.get('value', [])]
                
        except Exception as e:
            logger.error(f"Erro ao listar arquivos recentes: {e}")
            return []
    
    async def download_and_get_info(self, file_id: str) -> Optional[dict]:
        """
        Baixa um arquivo do OneDrive e retorna seus dados junto com o conteúdo.
        
        Combina download_file e get_file_info em uma única operação eficiente.
        
        Args:
            file_id: ID do arquivo no OneDrive
        
        Returns:
            dict com:
            - content: bytes do arquivo
            - info: metadados do arquivo (nome, tamanho, etc)
            Ou None se falhar
        
        Requirements: 7.4
        """
        headers = await self._get_headers()
        if not headers:
            logger.warning("OneDrive não conectado - não é possível baixar arquivo")
            return None
        
        try:
            async with httpx.AsyncClient() as client:
                # Obtém informações do arquivo incluindo URL de download
                response = await client.get(
                    f"{GRAPH_API_URL}/me/drive/items/{file_id}",
                    headers=headers,
                    params={'$select': 'id,name,size,webUrl,createdDateTime,lastModifiedDateTime,folder,file,parentReference,@microsoft.graph.downloadUrl'},
                    timeout=10.0
                )
                
                if response.status_code != 200:
                    logger.error(f"Erro ao obter arquivo: {response.status_code}")
                    return None
                
                item = response.json()
                download_url = item.get('@microsoft.graph.downloadUrl')
                
                # Se não tem URL de download direta, usa o endpoint /content
                if not download_url:
                    logger.info("URL de download não disponível, usando endpoint /content")
                    download_response = await client.get(
                        f"{GRAPH_API_URL}/me/drive/items/{file_id}/content",
                        headers=headers,
                        timeout=120.0,
                        follow_redirects=True
                    )
                else:
                    # Baixa o arquivo usando a URL direta
                    download_response = await client.get(
                        download_url,
                        timeout=120.0
                    )
                
                if download_response.status_code != 200:
                    logger.error(f"Erro no download: {download_response.status_code}")
                    return None
                
                file_info = self._parse_file_item(item)
                logger.info(f"Arquivo '{file_info.get('name')}' baixado com sucesso ({file_info.get('size_formatted')})")
                
                return {
                    'content': download_response.content,
                    'info': file_info
                }
                
        except httpx.TimeoutException:
            logger.error("Timeout ao baixar arquivo do OneDrive")
            return None
        except Exception as e:
            logger.error(f"Erro ao baixar arquivo: {e}")
            return None


async def send_onedrive_file_to_telegram(
    bot,
    file_id: str,
    chat_id: int,
    topic_id: int = None,
    caption: str = None,
    classify: bool = True
) -> Optional[dict]:
    """
    Baixa arquivo do OneDrive e envia para o Telegram.
    
    Fluxo completo:
    1. Baixa arquivo do OneDrive
    2. Classifica o documento (se classify=True)
    3. Envia para o tópico correto do Telegram
    4. Salva referência no banco de dados
    
    Args:
        bot: Instância do bot do Telegram
        file_id: ID do arquivo no OneDrive
        chat_id: ID do chat/grupo do Telegram
        topic_id: ID do tópico para enviar (se None, usa classificação automática)
        caption: Legenda para o arquivo (se None, usa nome do arquivo)
        classify: Se True, classifica o documento e envia para tópico correto
    
    Returns:
        dict com informações do documento salvo ou None se falhar:
        - telegram_file_id: file_id do Telegram
        - message_id: ID da mensagem enviada
        - topic_id: ID do tópico onde foi enviado
        - categoria: Categoria do documento
        - documento: Registro salvo no banco
    
    Requirements: 7.4
    """
    from config import CATEGORIA_TOPICO
    import database_adapter as db
    import ai
    import io
    
    # Baixa o arquivo do OneDrive
    result = await onedrive_client.download_and_get_info(file_id)
    
    if not result:
        logger.error(f"Falha ao baixar arquivo {file_id} do OneDrive")
        return None
    
    content = result['content']
    file_info = result['info']
    file_name = file_info.get('name', 'arquivo')
    file_extension = file_info.get('extension', '').lower()
    
    # Determina categoria e tópico
    categoria = 'outros'
    target_topic_id = topic_id
    
    if classify and not topic_id:
        # Classifica o documento usando IA
        try:
            # Para imagens, usa análise de imagem
            if file_extension in ['jpg', 'jpeg', 'png', 'gif', 'bmp', 'webp']:
                analise = await ai.analyze_image(content, caption or file_name)
                categoria = await ai.classify_document(str(analise))
            else:
                # Para outros arquivos, classifica pelo nome e extensão
                categoria = await ai.classify_document(f"Arquivo: {file_name}")
            
            target_topic_id = CATEGORIA_TOPICO.get(categoria, CATEGORIA_TOPICO['outros'])
            
        except Exception as e:
            logger.error(f"Erro ao classificar documento: {e}")
            categoria = 'outros'
            target_topic_id = CATEGORIA_TOPICO['outros']
    
    if not target_topic_id:
        target_topic_id = CATEGORIA_TOPICO.get('outros', 8)
    
    # Prepara legenda
    final_caption = caption or f"📄 {file_name}"
    if len(final_caption) > 200:
        final_caption = final_caption[:197] + "..."
    
    # Envia para o Telegram
    try:
        # Cria objeto de arquivo em memória
        file_bytes = io.BytesIO(content)
        file_bytes.name = file_name
        
        # Determina se é imagem ou documento
        is_image = file_extension in ['jpg', 'jpeg', 'png', 'gif', 'bmp', 'webp']
        
        if is_image:
            sent_msg = await bot.send_photo(
                chat_id=chat_id,
                photo=file_bytes,
                caption=final_caption,
                message_thread_id=target_topic_id
            )
            telegram_file_id = sent_msg.photo[-1].file_id if sent_msg.photo else None
        else:
            sent_msg = await bot.send_document(
                chat_id=chat_id,
                document=file_bytes,
                filename=file_name,
                caption=final_caption,
                message_thread_id=target_topic_id
            )
            telegram_file_id = sent_msg.document.file_id if sent_msg.document else None
        
        # Salva referência no banco
        descricao = caption or file_name
        if len(descricao) > 100:
            descricao = descricao[:97] + "..."
        
        documento = db.add_documento(
            tipo=file_extension or 'documento',
            descricao=descricao,
            file_id=telegram_file_id,
            categoria=categoria,
            message_id=sent_msg.message_id,
            topic_id=target_topic_id,
            dados_extraidos={
                'onedrive_file_id': file_id,
                'onedrive_name': file_name,
                'onedrive_size': file_info.get('size'),
                'onedrive_path': file_info.get('path'),
                'source': 'onedrive'
            }
        )
        
        logger.info(f"Arquivo '{file_name}' enviado para Telegram (tópico {target_topic_id}, categoria {categoria})")
        
        return {
            'telegram_file_id': telegram_file_id,
            'message_id': sent_msg.message_id,
            'topic_id': target_topic_id,
            'categoria': categoria,
            'documento': documento,
            'file_name': file_name,
            'file_info': file_info
        }
        
    except Exception as e:
        logger.error(f"Erro ao enviar arquivo para Telegram: {e}")
        return None


async def search_and_send_from_onedrive(
    bot,
    query: str,
    chat_id: int,
    reply_topic_id: int = None,
    limit: int = 1
) -> List[dict]:
    """
    Busca arquivos no OneDrive e envia para o Telegram.
    
    Função de conveniência que combina busca e envio.
    
    Args:
        bot: Instância do bot do Telegram
        query: Termo de busca
        chat_id: ID do chat/grupo do Telegram
        reply_topic_id: Tópico para responder (se None, classifica automaticamente)
        limit: Número máximo de arquivos para enviar
    
    Returns:
        Lista de resultados do envio (um dict por arquivo enviado)
    
    Requirements: 7.1, 7.4
    """
    # Busca arquivos no OneDrive
    files = await onedrive_client.search_files(query, limit=limit)
    
    if not files:
        logger.info(f"Nenhum arquivo encontrado no OneDrive para '{query}'")
        return []
    
    results = []
    
    for file_info in files:
        # Pula pastas
        if file_info.get('is_folder'):
            continue
        
        file_id = file_info.get('id')
        if not file_id:
            continue
        
        # Envia para o Telegram
        result = await send_onedrive_file_to_telegram(
            bot=bot,
            file_id=file_id,
            chat_id=chat_id,
            topic_id=reply_topic_id,
            caption=f"📁 Do OneDrive: {file_info.get('name')}",
            classify=True
        )
        
        if result:
            results.append(result)
    
    return results


async def smart_search_onedrive(query: str, max_files_to_read: int = 5) -> dict:
    """
    Busca inteligente no OneDrive.
    
    Estratégia:
    1. Primeiro busca pelo nome do arquivo
    2. Se não encontrar, lista arquivos recentes e lê o conteúdo
    3. Usa IA para encontrar o arquivo que corresponde à busca
    
    Args:
        query: O que o usuário está procurando (ex: "contrato do João")
        max_files_to_read: Máximo de arquivos para ler conteúdo
    
    Returns:
        dict com:
        - found: bool - se encontrou algo
        - method: 'name' ou 'content' - como encontrou
        - files: lista de arquivos encontrados
        - message: mensagem para o usuário
    """
    import ai
    import pdf_reader
    
    # 1. Primeiro tenta buscar pelo nome
    files_by_name = await onedrive_client.search_files(query, limit=5)
    
    # Filtra apenas arquivos (não pastas)
    files_by_name = [f for f in files_by_name if not f.get('is_folder')]
    
    if files_by_name:
        logger.info(f"Busca inteligente: encontrou {len(files_by_name)} arquivo(s) pelo nome")
        return {
            'found': True,
            'method': 'name',
            'files': files_by_name,
            'message': f"Achei {len(files_by_name)} arquivo(s) pelo nome!"
        }
    
    # 2. Não encontrou pelo nome - busca por conteúdo
    logger.info(f"Busca inteligente: não encontrou pelo nome, buscando por conteúdo...")
    
    # Lista arquivos recentes (PDFs e DOCs são os mais prováveis de ter o conteúdo)
    recent_files = await onedrive_client.get_recent_files(limit=20)
    
    # Filtra apenas arquivos legíveis (PDF, DOC, DOCX, TXT)
    readable_extensions = ['pdf', 'doc', 'docx', 'txt', 'xlsx', 'xls']
    readable_files = [
        f for f in recent_files 
        if not f.get('is_folder') and f.get('extension', '').lower() in readable_extensions
    ]
    
    if not readable_files:
        logger.info("Busca inteligente: nenhum arquivo legível encontrado")
        return {
            'found': False,
            'method': 'content',
            'files': [],
            'message': "Não encontrei arquivos para ler. Tenta outro termo!"
        }
    
    # Limita quantidade de arquivos para ler
    files_to_read = readable_files[:max_files_to_read]
    
    logger.info(f"Busca inteligente: lendo conteúdo de {len(files_to_read)} arquivo(s)...")
    
    # 3. Lê o conteúdo de cada arquivo e verifica se corresponde à busca
    matching_files = []
    
    for file_info in files_to_read:
        file_id = file_info.get('id')
        file_name = file_info.get('name', '')
        extension = file_info.get('extension', '').lower()
        
        try:
            # Baixa o arquivo
            result = await onedrive_client.download_and_get_info(file_id)
            
            if not result:
                continue
            
            content_bytes = result['content']
            content_text = ""
            
            # Extrai texto do arquivo
            if extension == 'pdf':
                # Usa o pdf_reader para extrair texto
                # extract_text_from_pdf retorna (text, has_text)
                text, has_text = pdf_reader.extract_text_from_pdf(content_bytes)
                if has_text:
                    content_text = text[:2000]  # Limita para não sobrecarregar
            elif extension == 'txt':
                try:
                    content_text = content_bytes.decode('utf-8', errors='ignore')[:2000]
                except:
                    content_text = ""
            elif extension in ['doc', 'docx']:
                # Para DOC/DOCX, usa python-docx se disponível
                try:
                    from docx import Document
                    from io import BytesIO
                    doc = Document(BytesIO(content_bytes))
                    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                    content_text = '\n'.join(paragraphs)[:2000]
                except ImportError:
                    # python-docx não instalado, tenta extrair texto básico
                    try:
                        content_text = content_bytes.decode('utf-8', errors='ignore')[:2000]
                    except:
                        content_text = ""
                except Exception as e:
                    logger.warning(f"Erro ao ler DOCX {file_name}: {e}")
                    content_text = ""
            elif extension in ['xlsx', 'xls']:
                # Para Excel, tenta extrair texto das células
                try:
                    import openpyxl
                    from io import BytesIO
                    wb = openpyxl.load_workbook(BytesIO(content_bytes), read_only=True)
                    texts = []
                    for sheet in wb.worksheets[:3]:  # Limita a 3 planilhas
                        for row in sheet.iter_rows(max_row=50, values_only=True):
                            row_text = ' '.join(str(cell) for cell in row if cell)
                            if row_text.strip():
                                texts.append(row_text)
                    content_text = '\n'.join(texts)[:2000]
                    wb.close()
                except ImportError:
                    content_text = ""
                except Exception as e:
                    logger.warning(f"Erro ao ler Excel {file_name}: {e}")
                    content_text = ""
            
            if not content_text:
                continue
            
            # Usa IA para verificar se o conteúdo corresponde à busca
            prompt = f"""Analise se este documento corresponde à busca do usuário.

BUSCA DO USUÁRIO: "{query}"

NOME DO ARQUIVO: {file_name}

CONTEÚDO DO DOCUMENTO (primeiros 2000 caracteres):
{content_text}

Responda APENAS com "SIM" se o documento corresponde à busca, ou "NAO" se não corresponde.
Considere correspondência se:
- O conteúdo menciona o assunto da busca
- O documento é sobre o tema buscado
- Há informações relevantes para a busca"""

            response = ai.model.generate_content(prompt)
            answer = response.text.strip().upper()
            
            if "SIM" in answer:
                matching_files.append(file_info)
                logger.info(f"Busca inteligente: arquivo '{file_name}' corresponde à busca!")
            
        except Exception as e:
            logger.error(f"Erro ao ler arquivo {file_name}: {e}")
            continue
    
    if matching_files:
        return {
            'found': True,
            'method': 'content',
            'files': matching_files,
            'message': f"Encontrei {len(matching_files)} arquivo(s) lendo o conteúdo! 📖"
        }
    
    return {
        'found': False,
        'method': 'content',
        'files': [],
        'message': f"Li {len(files_to_read)} arquivo(s) mas não encontrei nada sobre '{query}'"
    }


# Instância global para uso no bot
onedrive_auth = OneDriveAuth()
onedrive_client = OneDriveClient()
